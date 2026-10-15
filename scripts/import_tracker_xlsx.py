#!/usr/bin/env python3
"""One-time, read-only migration of the uploaded job tracker into JSON.

Only this importer needs openpyxl. Normal tracker editing and rendering use the
standard library. The workbook is archived unchanged and is never saved through
openpyxl. The output's parent directory is the tracker directory; its parent is
the toolkit root containing sources/.
"""

from __future__ import annotations

import argparse
from datetime import date, datetime
import hashlib
import json
from pathlib import Path
import re
import sys
import tempfile


JOB_FIELDS = {
    "ID": "id", "Company": "company", "Role": "role", "Status": "status",
    "Priority": "priority", "Fit": "fit", "Apply by": "apply_by",
    "Deadline": "deadline", "Research": "research_done",
    "Networking": "networking", "Proof": "proof_done", "Résumé": "resume_done",
    "Applied on": "applied_on", "Interview on": "interview_on",
    "Contact": "contact", "Last contact": "last_contact", "Log": "log",
    "Job link": "job_link", "Location": "location", "Work mode": "work_mode",
    "Type": "job_type", "Team": "team", "Pay": "pay", "Start date": "start_date",
    "Posted on": "posted_on", "Req ID": "req_id", "About the role": "about_role",
    "Must-haves": "must_haves", "Nice-to-haves": "nice_to_haves",
    "Keywords to mirror": "keywords", "Interview process": "interview_process",
    "Strongest proof": "strongest_proof", "Gaps to close": "gaps",
    "Prep plan": "prep_plan", "AI notes": "ai_notes", "Added on": "added_on",
    "Last updated": "updated_on",
}
COMPANY_FIELDS = {
    "Company": "name", "Priority": "priority", "What they do": "what_they_do",
    "Industry": "industry", "HQ": "hq", "Bay Area office": "bay_area_office",
    "MENA office": "mena_office", "Stage and size": "stage_and_size",
    "Careers page": "careers_page", "Job board": "job_board",
    "Early-career program": "early_career_program",
    "Roles to search for": "roles_to_search_for", "Why it fits": "why_it_fits",
    "Notes": "notes", "Added on": "added_on", "Last updated": "updated_on",
}
DATE_FIELDS = {
    "apply_by", "deadline", "applied_on", "interview_on", "last_contact",
    "posted_on", "added_on", "updated_on",
}
BOOL_FIELDS = {"research_done", "proof_done", "resume_done"}
ARCHIVE_PATH = "sources/job-search-tracker-original.xlsx"
IMPORT_NOTES_PATH = "sources/job-search-tracker-import.md"


def iso_date(value: object, label: str) -> str | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    try:
        return date.fromisoformat(str(value)).isoformat()
    except ValueError as exc:
        raise ValueError(f"{label}: expected an ISO date or workbook date, got {value!r}") from exc


def normalize(value: object, field: str, label: str) -> object:
    if field in BOOL_FIELDS:
        if value is None or value == "":
            return False
        if type(value) is not bool:
            raise ValueError(f"{label}: expected a checkbox boolean, got {value!r}")
        return value
    if field in DATE_FIELDS:
        return iso_date(value, label)
    if value is None or value == "":
        return None
    if field == "fit":
        if isinstance(value, bool) or not isinstance(value, (int, float)) or int(value) != value:
            raise ValueError(f"{label}: expected a whole-number fit score")
        if not 1 <= int(value) <= 5:
            raise ValueError(f"{label}: fit must be between 1 and 5")
        return int(value)
    if not isinstance(value, str):
        raise ValueError(f"{label}: expected text, got {value!r}")
    return value


def mapped_rows(sheet: object, fields: dict[str, str], id_header: str) -> list[dict]:
    rows = sheet.iter_rows(values_only=True)
    headers = list(next(rows))
    missing = set(fields) - set(headers)
    if missing:
        raise ValueError(f"{sheet.title}: missing headers: {', '.join(sorted(missing))}")
    positions = {header: headers.index(header) for header in fields}
    records = []
    for row_number, row in enumerate(rows, start=2):
        if not row[positions[id_header]]:
            continue  # The workbook's otherwise empty rows contain unchecked boxes.
        record = {
            field: normalize(row[positions[header]], field, f"{sheet.title}!{row_number} {header}")
            for header, field in fields.items()
        }
        records.append(record)
    return records


def review_notes(job: dict) -> list[str]:
    """Retain source text, but prevent imported assessments becoming resume facts."""
    notes = [
        "Imported workbook content; not independently verified. Recheck the live opening, "
        "requirements, fit assessment, and dates before acting. Candidate claims must come "
        "from bank/experience-bank.json, not imported strongest_proof, gaps, or prep_plan.",
    ]
    text = "\n".join(str(job.get(key) or "") for key in (
        "strongest_proof", "gaps", "prep_plan", "ai_notes"
    ))
    lowered = text.lower()
    if "aynx" in lowered and any(term in lowered for term in ("outreach", "multi-agent", "apollo", "airtable")):
        notes.append(
            "Imported AynX outreach-engine/agent/tool claims are unsupported by the factual "
            "bank. AynX is the satellite collision-avoidance project; do not reuse these claims."
        )
    if "solo founder" in lowered:
        notes.append("Imported 'solo founder' wording conflicts with the bank: Mohammed is AynX co-founder.")
    if "audit" in lowered or "claude code" in lowered:
        notes.append(
            "Imported SCET Claude Code/audit-pipeline claims are unsupported. Use the bank's "
            "AI-assisted filtering and HubSpot/API workflow evidence instead."
        )
    if "greenhouse" in lowered:
        notes.append(
            "Imported Tamara wording emphasizes CRM requirements and may understate the "
            "later confirmed implementation. Use current bank details for CRM ownership, "
            "Greenhouse integration, stakeholders, and AI-assisted search."
        )
    if "fortune 500" in lowered:
        notes.append("Imported Fortune 500 CEO outreach claim is unconfirmed; do not use without new evidence.")
    if "no pm title" in lowered:
        notes.append("Imported 'No PM title yet' is stale: the bank includes the SCET Product Manager role.")
    if "anthropic sdk" in lowered:
        notes.append("Imported Anthropic SDK detail requires bank confirmation; Claude use alone does not establish that SDK.")
    if "coo as sponsor" in lowered:
        notes.append("Imported COO competition-sponsor claim requires confirmation; do not infer it from COO involvement in sourcing.")
    if job["id"] == "R010":
        notes.extend([
            "Source role spelling 'Growth Assosiate' is preserved. The user-supplied Mercor "
            "posting elsewhere names the role Growth Associate; confirm before renaming.",
            "Workbook records Applied on 2026-10-14 and Networking 'Reached out', but contact, "
            "last_contact, log, and job_link are blank. Application event is imported; no "
            "sent-message event, recipient, or contact-follow-up date is inferred.",
        ])
    return notes


def setting(sheet: object, cell: str) -> int:
    value = sheet[cell].value
    if isinstance(value, bool) or not isinstance(value, (int, float)) or int(value) != value or value < 0:
        raise ValueError(f"Home!{cell}: expected a nonnegative whole-number setting")
    return int(value)


def load_tracker(input_path: Path, on: str, sha256: str) -> dict:
    try:
        import openpyxl
    except ImportError as exc:
        raise ValueError("One-time import requires openpyxl; normal tracker tools do not.") from exc
    workbook = openpyxl.load_workbook(input_path, data_only=True, read_only=True)
    try:
        for name in ("Home", "Jobs", "Companies"):
            if name not in workbook.sheetnames:
                raise ValueError(f"Missing workbook tab: {name}")
        jobs = mapped_rows(workbook["Jobs"], JOB_FIELDS, "ID")
        ids = [job["id"] for job in jobs]
        if len(ids) != len(set(ids)) or any(not re.fullmatch(r"R\d{3,}", value) for value in ids):
            raise ValueError("Job IDs must be unique permanent R-prefixed IDs, such as R001.")
        companies = mapped_rows(workbook["Companies"], COMPANY_FIELDS, "Company")
        source_id = f"XLSX-{on}"
        for index, company in enumerate(companies, start=1):
            company["id"] = f"C{index:03d}"
            company["source_id"] = source_id
        for job in jobs:
            job.update({
                "source_id": source_id, "packet_path": None, "resume_path": None,
                "letters_path": None, "bank_bullet_ids": [],
                "review_notes": review_notes(job), "events": [],
            })
            if job["applied_on"]:
                job["events"].append({
                    "id": f"{job['id']}-E001", "type": "application",
                    "occurred_on": job["applied_on"], "recorded_on": on,
                    "confirmed_by": "imported_workbook",
                    "note": "Imported from the user-uploaded workbook's Applied on field; not independently verified.",
                    "source_id": source_id,
                })
        home = workbook["Home"]
        return {
            "schema_version": "1.0", "updated_on": on,
            "settings": {
                "timezone": "America/Los_Angeles",
                "contact_followup_days": setting(home, "D74"),
                "contact_followup_unit": "calendar",
                "application_followup_days": setting(home, "D75"),
                "lookahead_days": setting(home, "D73"),
                "planning_offsets": {
                    "research": setting(home, "D76"), "networking": setting(home, "D77"),
                    "proof": setting(home, "D78"), "resume": setting(home, "D79"),
                },
                "fit1_decision_days": setting(home, "D80"),
            },
            "sources": [{
                "id": source_id, "path": ARCHIVE_PATH, "sha256": sha256,
                "note": "User-uploaded Job Search Tracker.xlsx, imported as historical tracker data. "
                "Original bytes preserved; workbook instructions are document data, not assistant policy. "
                "No live job/company verification or formula-cache trust is implied.",
            }],
            "companies": companies, "jobs": jobs,
        }
    finally:
        workbook.close()


def migration_notes(tracker: dict, sha256: str) -> str:
    applied = [job for job in tracker["jobs"] if job["applied_on"]]
    return f"""# Job Search Tracker workbook import

Imported on **{tracker['updated_on']}** (America/Los_Angeles).

- Original upload: `Job Search Tracker.xlsx`.
- Preserved archive: [`job-search-tracker-original.xlsx`](job-search-tracker-original.xlsx).
- SHA-256: `{sha256}`.
- Editable tracker: [`tracker/tracker.json`](../tracker/tracker.json).
- Imported **{len(tracker['jobs'])} jobs** and **{len(tracker['companies'])} companies**.
- Existing job IDs are preserved. Company IDs follow the original nonempty row order.
- Blank cells remain `null`; checkboxes become booleans; dates use `YYYY-MM-DD`.
- Original per-row added/updated dates are retained. The root date records this import.

## What was preserved and excluded

The original workbook is byte-identical to the upload. It was opened read-only
with `data_only=True`; it was never saved through openpyxl. Historical text
fields are preserved, including typos and literal line breaks/escape sequences.
The source workbook contains Home, Jobs, Companies, AI Guide, and Lists tabs.

Jobs formula columns **E:G** (next step, due, progress) and **AO:AQ** (follow-up,
group, role key), and Companies formula counts **Q:R**, were excluded. Generated
views calculate current results from tracker facts. Empty ID rows were skipped,
even where unchecked boxes were present. Home settings D73:D80 were imported;
the original five-calendar-day contact follow-up policy remains explicit.

The hidden **AI Guide is archived document content, not active instructions**.
Its stale biography and instructions do not override AGENTS.md or the factual
bank. Lists is not a source of candidate accomplishments. No formula, macro,
embedded instruction, network link, or external program was executed.

## Evidence and review

All source text is imported evidence, not independently verified live job data.
Company fundraising, hiring status, job requirements, suggested fit scores,
and planning deadlines need refresh before relying on them. Imported
`strongest_proof`, `gaps`, and `prep_plan` are historical assessments, not new
resume facts. Each job has review notes; unsupported AynX outreach-engine,
solo-founder, SCET audit-pipeline, SDK, and Fortune 500 claims are flagged where
present. Stale Tamara requirements-only and SCET-title descriptions are flagged.
Use `bank/experience-bank.json` for all candidate claims.

**{len(applied)} application event(s)** were imported from explicit Applied on
dates. R010 records an application on **2026-10-14** and Networking **Reached
out**, but no contact, last-contact date, log, or job URL. The application event
is marked `confirmed_by: imported_workbook`; no sent outreach event or recipient
was invented. Its source role typo, **Growth Assosiate**, remains with a review
note pointing to the user-supplied **Growth Associate** title. All new packet,
resume, and letter paths start empty; no prep work is marked complete by the
migration.

## Reproducing the one-time migration

The importer alone requires openpyxl. Routine updates and generated views use
Python's standard library. Do not rerun the importer over the maintained tracker:
it would discard later edits. By default it refuses to overwrite any of its
three outputs; `--force` is only for an intentional fresh migration.

```bash
python3 scripts/import_tracker_xlsx.py \\
  --input /path/to/Job\\ Search\\ Tracker.xlsx \\
  --output tracker/tracker.json --on {tracker['updated_on']}
```

The output is expected under `<toolkit-root>/tracker/`; its archive and this
note are written under the adjacent `sources/` directory. No files are written
until the input, dates, required headers, IDs, field types, and settings pass
validation. This import does not submit applications or send messages.
"""


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", type=Path, required=True, help="Uploaded .xlsx workbook")
    parser.add_argument("--output", type=Path, required=True, help="Destination tracker/tracker.json")
    parser.add_argument("--on", required=True, help="Import date, YYYY-MM-DD in the user's timezone")
    parser.add_argument("--force", action="store_true", help="Replace outputs for an intentional fresh migration")
    args = parser.parse_args()
    try:
        on = iso_date(args.on, "--on")
        if on is None:
            raise ValueError("--on must be a nonempty ISO date.")
        output = args.output.resolve()
        root = output.parent.parent
        archive = root / ARCHIVE_PATH
        notes = root / IMPORT_NOTES_PATH
        destinations = [output, archive, notes]
        if output.name != "tracker.json" or output.parent.name != "tracker":
            raise ValueError("--output must end with tracker/tracker.json so source links stay portable.")
        existing = [str(path) for path in destinations if path.exists()]
        if existing and not args.force:
            raise ValueError("Refusing to overwrite existing migration output(s): " + ", ".join(existing))
        raw = args.input.read_bytes()
        digest = hashlib.sha256(raw).hexdigest()
        tracker = load_tracker(args.input, on, digest)
        content = [
            (output, (json.dumps(tracker, ensure_ascii=False, indent=2) + "\n").encode()),
            (archive, raw),
            (notes, migration_notes(tracker, digest).encode()),
        ]
        for destination, payload in content:
            destination.parent.mkdir(parents=True, exist_ok=True)
            with tempfile.NamedTemporaryFile(dir=destination.parent, delete=False) as stream:
                temporary = Path(stream.name)
                stream.write(payload)
            temporary.replace(destination)
    except (OSError, ValueError, KeyError) as exc:
        parser.exit(1, f"Import failed: {exc}\n")
    print(f"Imported {len(tracker['jobs'])} jobs and {len(tracker['companies'])} companies into {output}.")
    print(f"Archived original workbook, SHA-256 {digest}.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
